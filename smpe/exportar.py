"""Exportacao das tabelas do sistema para Excel (.xlsx), sempre no mesmo layout."""
import io
import re
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .util import fmt_cpf, sem_acento

HORA_LOCAL = timezone(timedelta(hours=-3))  # Sao Luis (sem horario de verao)

NAVY, AZUL, ZEBRA, LINHA, CINZA = "0A2657", "0B63CE", "F3F7FC", "D5DFEC", "6B7A93"
FONTE = "Calibri"
BORDA = Border(*(Side(style="thin", color=LINHA),) * 4)
FORMATOS = {"numero": "#,##0", "decimal": "#,##0.00", "moeda": '"R$" #,##0.00', "pct": "0.0%",
            "data": "DD/MM/YYYY", "datahora": "DD/MM/YYYY HH:MM"}
TIPOS = {"texto", "cpf", *FORMATOS}


def _valor(v, tipo: str):
    """Converte o valor da tela para o tipo da celula (numero, data...) ou texto."""
    if v is None or v == "":
        return None
    if tipo in ("numero", "decimal", "moeda", "pct"):
        try:
            return float(v) if tipo != "numero" else int(float(v))
        except (TypeError, ValueError):
            return str(v)
    if tipo in ("data", "datahora"):
        s = str(v)
        try:
            d = datetime.fromisoformat(s.replace("T", " ")[:19])
        except ValueError:
            return s
        return d if tipo == "datahora" else d.date()
    if tipo == "cpf":
        s = str(v)
        return fmt_cpf(s) if s.isdigit() else s
    return str(v)


def _largura(v) -> int:
    if isinstance(v, datetime):
        return 16
    if isinstance(v, date):
        return 11
    if isinstance(v, float):
        return len(f"{v:,.2f}")
    return max((len(p) for p in str(v).split("\n")), default=0)


def nome_arquivo(titulo: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", sem_acento(titulo)).strip("_").upper() or "PLANILHA"
    return f"{base[:60]}_{datetime.now(HORA_LOCAL):%Y-%m-%d}.xlsx"


def gerar(titulo: str, colunas: list[dict], linhas: list[list], subtitulo: str = "", usuario: str = "") -> bytes:
    """colunas: [{"titulo": str, "tipo": texto|numero|decimal|moeda|pct|data|datahora|cpf}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", titulo)[:31] or "Planilha"
    n = max(len(colunas), 1)
    tipos = [c.get("tipo") if c.get("tipo") in TIPOS else "texto" for c in colunas]

    # titulo e identificacao
    ws.cell(1, 1, titulo).font = Font(name=FONTE, size=14, bold=True, color=NAVY)
    agora = datetime.now(HORA_LOCAL)
    info = " · ".join(x for x in ("SIS SMPE", subtitulo, f"Gerado em {agora:%d/%m/%Y %H:%M}"
                                  + (f" por {fmt_cpf(usuario)}" if usuario else "")) if x)
    ws.cell(2, 1, info).font = Font(name=FONTE, size=9, italic=True, color=CINZA)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.row_dimensions[1].height = 22

    # cabecalho
    H = 4
    for j, c in enumerate(colunas, 1):
        cel = ws.cell(H, j, c.get("titulo", ""))
        cel.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=NAVY)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA
    ws.row_dimensions[H].height = 24

    # linhas (zebradas, com bordas finas)
    larg = [len(str(c.get("titulo", ""))) for c in colunas]
    zebra = PatternFill("solid", fgColor=ZEBRA)
    for i, lin in enumerate(linhas, 1):
        for j, tipo in enumerate(tipos, 1):
            v = _valor(lin[j - 1] if j - 1 < len(lin) else None, tipo)
            cel = ws.cell(H + i, j, v)
            cel.font = Font(name=FONTE, size=10)
            cel.border = BORDA
            num = tipo in FORMATOS and not isinstance(v, str)
            if num:
                cel.number_format = FORMATOS[tipo]
            cel.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and "\n" in v,
                                      horizontal="right" if num and tipo not in ("data", "datahora") else
                                      "center" if tipo in ("data", "datahora", "cpf") else "left")
            if i % 2 == 0:
                cel.fill = zebra
            if v is not None:
                larg[j - 1] = max(larg[j - 1], _largura(v))

    # total de registros
    fim = H + len(linhas)
    tot = ws.cell(fim + 2, 1, f"Total: {len(linhas)} registro(s)")
    tot.font = Font(name=FONTE, size=10, bold=True, color=AZUL)

    for j, w in enumerate(larg, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 9), 60)
    ws.freeze_panes = ws.cell(H + 1, 1)
    if colunas:
        ws.auto_filter.ref = f"A{H}:{get_column_letter(n)}{max(fim, H)}"
    ws.print_title_rows = f"{H}:{H}"
    ws.page_setup.orientation = "landscape" if n > 5 else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Página &P de &N"
    ws.oddFooter.right.text = "Uso restrito (LGPD)"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
