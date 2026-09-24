"""Importacao das bases (planilha SIS SMPE completa ou arquivos avulsos .xlsx/.xlsm/.csv)."""
import csv
import io
import warnings
from pathlib import Path

import openpyxl

from . import db
from .util import norm_cpf, norm_nome, sem_acento, to_date, txt

warnings.filterwarnings("ignore", module="openpyxl")

# base -> (aba padrao na planilha, coluna obrigatoria, {coluna_db: [cabecalhos aceitos]})
BASES = {
    "escolas": ("CAD_ESCOLA", "ESCOLA", {
        "id": ["ID.:", "ID"], "nome": ["ESCOLA"], "cod_smtt": ["COD_SMTT", "COD SMTT"],
        "inep": ["CODIGO DO INEP", "INEP"]}),
    "geduc": ("GEDUC", "ALUNO", {
        "inep_escola": ["INEP_ESCOLA"], "escola": ["ESCOLA"], "nucleo": ["NUCLEO"], "modalidade": ["MODALIDADE"],
        "id_ano_serie": ["ID_ANO_SERIE"], "aluno": ["ALUNO"], "mae": ["MAE"], "pai": ["PAI"], "genero": ["GENERO"],
        "ano_serie": ["ANO_SERIE"], "turno": ["TURNO"], "turma": ["TURMA"], "id_aluno": ["ID_ALUNO"],
        "dt_nasc": ["DT_NASCIMENTO"], "rua": ["RUA"], "numero": ["NUMERO"], "bairro": ["BAIRRO"],
        "cidade": ["CIDADE"], "cep": ["CEP"], "cpf": ["CPF_ALUNO", "CPF"]}),
    "censo": ("CENSO", "NOME DO ALUNO", {
        "id_inep": ["IDENTIFICACAO UNICA"], "aluno": ["NOME DO ALUNO"], "dt_nasc": ["DATA DE NASCIMENTO"],
        "cor": ["COR/RACA"], "sexo": ["SEXO"], "cpf": ["CPF"]}),
    "smtt": ("SMTT", "ESTUDANTE", {
        "portaria": ["PORTARIA"], "escola": ["REVISADO"], "id_smtt": ["ID"], "grau": ["GRAU"], "periodo": ["PERIODO"],
        "turma": ["TURMA"], "matricula": ["MATRICULA"], "cartao": ["CARTAO"], "estudante": ["ESTUDANTE"],
        "nascido": ["NASCIDO"], "cpf": ["CPF"], "rg": ["RG"], "org_exp": ["ORG EXP"], "data_exp": ["DATA EXP"],
        "sexo": ["SEXO"], "telefone": ["TELEFONE"], "celular": ["CELULAR"], "email": ["EMAIL"], "cep": ["CEP"],
        "endereco": ["ENDERECO"], "complemento": ["COMPLEMENTO"], "numero": ["NUMERO"], "bairro": ["BAIRRO"],
        "cidade": ["CIDADE"], "uf": ["UF"], "mae": ["MAE"], "pai": ["PAI"], "cadastrado": ["CADASTRADO"],
        "alterado": ["ALTERDADO", "ALTERADO"], "curso": ["CURSO"], "tipo": ["TIPO"], "modalidade": ["MODALIDADE"]}),
    "status_alunos": (" ALUNOS_POR_STATUS", "ALUNO", {
        "escola": ["ESCOLA"], "turma": ["TURMA"], "turno": ["TURNO"], "id_aluno": ["ID_ALUNO"],
        "inep_aluno": ["INEP_ALUNO"], "aluno": ["ALUNO"], "matricula": ["MATRICULA"], "nascimento": ["NASCIMENTO"],
        "idade": ["IDADE"], "sexo": ["SEXO"], "cor": ["COR"], "nis": ["NIS"], "situacao": ["SITUACAO"],
        "mae": ["NOME_MAE"], "pai": ["NOME_PAI"], "endereco": ["ENDERECO_ALUNO"], "numero": ["NUMERO"],
        "bairro": ["BAIRRO"], "certidao": ["CERTIDAO_NASCIMENTO"], "cpf": ["CPF_ALUNO"], "telefone": ["TELEFONE"]}),
}
DATE_COLS = {"dt_nasc", "nascido", "data_exp", "nascimento"}
CPF_COLS = {"cpf"}
NOME_COL = {"geduc": "aluno", "censo": "aluno", "smtt": "estudante", "status_alunos": "aluno"}
LABELS = {"escolas": "Cadastro de escolas", "geduc": "GEDUC", "censo": "Censo escolar", "smtt": "SMTT",
          "status_alunos": "Alunos por status"}


def _h(v) -> str:
    return sem_acento(txt(v)).upper().strip()


def _find_header(rows: list[tuple], obrigatoria: str) -> int:
    for i, r in enumerate(rows[:15]):
        if obrigatoria in {_h(c) for c in r}:
            return i
    raise ValueError(f"Cabecalho nao encontrado (coluna '{obrigatoria}')")


def _map_columns(header: tuple, cols: dict) -> dict[str, int]:
    hn = [_h(c) for c in header]
    out = {}
    for col, aliases in cols.items():
        for a in aliases:
            if a in hn:
                out[col] = hn.index(a)
                break
    return out


def _clean(base: str, col: str, v):
    if col in DATE_COLS:
        return to_date(v)
    if col in CPF_COLS:
        return norm_cpf(v)
    return txt(v)


def _rows_to_records(base: str, rows_iter) -> list[dict]:
    _, obrig, cols = BASES[base]
    head = []
    for r in rows_iter:
        head.append(r)
        if len(head) >= 15:
            break
    hi = _find_header(head, obrig)
    cmap = _map_columns(head[hi], cols)
    nome_idx = cmap.get(NOME_COL.get(base, "nome"))
    recs = []

    def conv(r):
        if nome_idx is None or nome_idx >= len(r) or not txt(r[nome_idx]):
            return None
        return {c: _clean(base, c, r[i] if i < len(r) else None) for c, i in cmap.items()}

    for r in head[hi + 1:]:
        x = conv(r)
        if x:
            recs.append(x)
    for r in rows_iter:
        x = conv(r)
        if x:
            recs.append(x)
    return recs


def read_file(base: str, filename: str, data: bytes, sheet: str | None = None) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        if text.count("�") > 5:
            text = data.decode("cp1252")
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=";,\t")
        return _rows_to_records(base, iter(list(csv.reader(io.StringIO(text), dialect))))
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        aba = sheet or BASES[base][0]
        ws = wb[aba] if aba in wb.sheetnames else _guess_sheet(wb, base)
        return _rows_to_records(base, ws.iter_rows(values_only=True))
    finally:
        wb.close()


def _guess_sheet(wb, base):
    alvo = BASES[base][0].strip().upper()
    for n in wb.sheetnames:
        if n.strip().upper() == alvo:
            return wb[n]
    return wb.worksheets[0]  # arquivo avulso com uma aba so


def save(con, base: str, recs: list[dict], arquivo: str) -> int:
    with con:
        if base == "escolas":
            for r in recs:
                if not str(r.get("id", "")).isdigit():
                    continue
                con.execute("""INSERT INTO escolas (id, nome, cod_smtt, inep) VALUES (?,?,?,?)
                               ON CONFLICT(id) DO UPDATE SET nome=excluded.nome, cod_smtt=excluded.cod_smtt,
                               inep=excluded.inep""",
                            (int(r["id"]), r["nome"], r.get("cod_smtt", ""), r.get("inep", "")))
        else:
            con.execute(f"DELETE FROM {base}")
            cols = list(BASES[base][2].keys())
            extra = ["nome_norm"] + (["escola_norm"] if base == "geduc" else [])
            sql = f"INSERT INTO {base} ({','.join(cols + extra)}) VALUES ({','.join('?' * (len(cols) + len(extra)))})"
            nome_col = NOME_COL[base]
            con.executemany(sql, [
                [r.get(c, "") for c in cols] + [norm_nome(r.get(nome_col))]
                + ([norm_nome(r.get("escola"))] if base == "geduc" else []) for r in recs])
        con.execute("INSERT INTO importacoes (base, arquivo, linhas) VALUES (?,?,?)", (base, arquivo, len(recs)))
    return len(recs)


def import_workbook(path: str, bases: list[str] | None = None, log=print) -> dict:
    """Importa a planilha SIS SMPE completa (todas as abas conhecidas)."""
    con = db.connect()
    data = Path(path).read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {}
    for base in bases or list(BASES):
        aba = BASES[base][0]
        if aba not in wb.sheetnames:
            log(f"{base}: aba '{aba}' nao encontrada, pulando")
            continue
        recs = _rows_to_records(base, wb[aba].iter_rows(values_only=True))
        out[base] = save(con, base, recs, Path(path).name)
        log(f"{LABELS[base]}: {out[base]} linhas")
    wb.close()
    return out
