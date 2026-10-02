"""Modulo Servidores (RH da SEMED): versao web da planilha SGI Servidor.xlsm.

Base: exportacao de funcionarios do GEDUC (ou a aba CADASTRO da propria planilha). A importacao atualiza o cadastro
pela matricula, sem apagar o que o RH completou no sistema (turno, quadro, status, atuacao...)."""
import re
from datetime import date, timedelta

from . import db, importer
from .util import norm_cpf, norm_nome, sem_acento, to_date, txt

# (aba padrao, coluna(s) obrigatoria(s), {coluna_db: [cabecalhos aceitos]}): GEDUC de funcionarios e aba CADASTRO do SGI
NOMES = ["PROFESSOR", "NOME DO FUNCIONARIO", "NOME DO SERVIDOR", "SERVIDOR", "FUNCIONARIO", "NOME"]
importer.OUTRAS["servidores"] = ("CADASTRO", NOMES, {
    "nome": NOMES, "cpf": ["CPF"], "sexo": ["SEXO", "GENERO"],
    "dt_nasc": ["NASCIMENTO", "DATA DE NASCIMENTO", "DT_NASCIMENTO", "DATA NASCIMENTO"], "matricula": ["MATRICULA"],
    "cargo": ["CARGO_CONCURSO", "CARGO"], "funcao": ["FUNCAO"], "lotacao": ["ESCOLA", "LOTACAO", "UNIDADE"],
    "zona": ["ZONA"], "setor": ["SETOR"], "carga_horaria": ["CARGA_HORARIA", "CARGA HORARIA"],
    "situacao_funcional": ["SITUACAO_FUNCIONAL", "SITUACAO FUNCIONAL"],
    "regime_contratacao": ["REGIME CONTRATACAO", "REGIME DE CONTRATACAO"],
    "turno": ["TURNO DE LOTACAO", "TURNO"], "horas_semanais": ["HORAS SEMANAIS"],
    "dt_admissao": ["DATA EXERCICIO: (ADMISSAO)", "DATA EXERCICIO", "DATA DE ADMISSAO", "ADMISSAO"],
    "quadro": ["QUADRO"], "regime_juridico": ["REGIME JURIDICO:", "REGIME JURIDICO"],
    "formacao": ["FORMACAO (GRADUACAO)", "FORMACAO"], "habilitacao": ["HABILITACAO (CURSO)", "HABILITACAO"],
    "status": ["STATUS"], "orgao": ["ORGAO"], "tipo_ensino": ["TIPO DE ENSINO"],
    "atuacao": ["ATUACAO DO PROFESSOR", "ATUACAO"], "componente": ["COMPONENTE CURRICULAR", "DISCIPLINA(S)", "DISCIPLINAS"]})

CAMPOS = ["nome", "cpf", "sexo", "dt_nasc", "matricula", "cargo", "funcao", "lotacao", "zona", "setor", "carga_horaria",
          "situacao_funcional", "regime_contratacao", "turno", "horas_semanais", "dt_admissao", "quadro",
          "regime_juridico", "formacao", "habilitacao", "status", "orgao", "tipo_ensino", "atuacao", "componente"]
SEM_MAIUSCULA = {"componente", "formacao", "habilitacao"}  # texto livre, como digitado
DATAS = {"dt_nasc", "dt_admissao"}

TURNOS = ["MATUTINO", "VESPERTINO", "NOTURNO", "INTEGRAL"]
QUADROS = ["MAGISTÉRIO", "TÉCNICO/ADMINISTRATIVO", "TERCEIRIZADO"]
STATUS = ["ATIVO", "INATIVO"]
TIPOS_ENSINO = ["Educação Infantil", "Ensino Fundamental I", "Ensino Fundamental II", "Educação de Jovens e Adultos"]
ATUACOES = ["Professor(a) Titular", "Professor(a) de Complementação"]
MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO",
         "NOVEMBRO", "DEZEMBRO"]


def _chave(v: str) -> str:
    return sem_acento(v).upper()


def normaliza(campo: str, v) -> str:
    """Valor do campo como fica no cadastro (padroniza as variacoes da planilha: 'MATTUTINO', 'TERCEIRAZADO'...)."""
    if campo in DATAS:
        return to_date(v)
    if campo == "cpf":
        return norm_cpf(v)
    s = re.sub(r"\s+", " ", txt(v)).strip()
    if re.fullmatch(r"Colunas?\d*", s, re.I) or s.upper() == "SEM CADASTRO NO GEDUC":  # sobras da planilha
        return ""
    if campo in ("matricula", "carga_horaria", "horas_semanais"):
        return re.sub(r"[^\dA-Za-z-]", "", s) if campo == "matricula" else s.upper().rstrip("H").strip()
    if campo in ("tipo_ensino", "atuacao"):  # grafia padrao das listas ('Professor(a) Títular' -> 'Titular')
        return next((x for x in (TIPOS_ENSINO if campo == "tipo_ensino" else ATUACOES) if _chave(x) == _chave(s)), s)
    if campo in SEM_MAIUSCULA:
        return s
    s, k = s.upper(), _chave(s)
    if campo == "sexo":
        return "FEMININO" if k.startswith("F") else "MASCULINO" if k.startswith("M") else s
    if campo == "turno":
        for t in TURNOS:
            if k.startswith(t[:3]) and "/" not in k:
                return t
    if campo == "quadro":
        for prefixo, q in (("MAGIST", QUADROS[0]), ("TEC", QUADROS[1]), ("TERC", QUADROS[2])):
            if k.startswith(prefixo):
                return q
    if campo == "status" and k in STATUS:
        return k
    return s


def limpa(r: dict) -> dict:
    return {c: normaliza(c, r.get(c)) for c in CAMPOS if c in r}


def importar(con, recs: list[dict], arquivo: str) -> dict:
    """Inclui os novos e atualiza os existentes pela matricula (sem matricula: pelo CPF ou nome). Campo vazio no
    arquivo nao apaga o que ja esta no cadastro; servidor que nao veio no arquivo continua cadastrado."""
    atuais = [dict(r) for r in con.execute("SELECT id, matricula, cpf, nome_norm FROM servidores")]
    por_mat = {r["matricula"]: r["id"] for r in atuais if r["matricula"]}
    sem_mat = {r["id"]: r for r in atuais if not r["matricula"]}  # inclui os criados neste mesmo arquivo
    com_mat = set(por_mat.values())
    por_cpf = {r["cpf"]: r["id"] for r in atuais if r["cpf"]}
    por_nome = {r["nome_norm"]: r["id"] for r in atuais}
    novos = atualizados = 0
    with con:
        for r in recs:
            d = {k: v for k, v in limpa(r).items() if v}
            if not d.get("nome"):
                continue
            nn = norm_nome(d["nome"])
            mat = d.get("matricula", "")
            if mat:  # mesmo servidor ja cadastrado sem matricula (cadastro manual) tambem e atualizado
                sid = por_mat.get(mat) or next((x["id"] for x in sem_mat.values() if x["cpf"] and x["cpf"] == d.get("cpf")
                                                or x["nome_norm"] == nn), None)
            else:
                sid = por_cpf.get(d.get("cpf", "")) or por_nome.get(nn)
            if sid:
                con.execute(f"UPDATE servidores SET {', '.join(f'{c}=?' for c in d)}, nome_norm=?, "
                            f"atualizado_em={_agora()} WHERE id=?", (*d.values(), nn, sid))
                atualizados += 1
            else:
                cols = list(d)
                sid = con.execute(f"INSERT INTO servidores ({', '.join(cols)}, nome_norm) "
                                  f"VALUES ({', '.join('?' * (len(cols) + 1))}) RETURNING id", (*d.values(), nn)).fetchone()[0]
                novos += 1
            if mat:
                por_mat[mat] = sid
                com_mat.add(sid)
                sem_mat.pop(sid, None)
            elif sid not in sem_mat and sid not in com_mat:
                sem_mat[sid] = {"id": sid, "cpf": d.get("cpf", ""), "nome_norm": nn}
            if d.get("cpf"):
                por_cpf[d["cpf"]] = sid
            por_nome[nn] = sid
        con.execute("INSERT INTO importacoes (base, arquivo, linhas) VALUES ('servidores',?,?)", (arquivo, len(recs)))
    total = con.execute("SELECT COUNT(*) FROM servidores").fetchone()[0]
    return {"lidos": len(recs), "novos": novos, "atualizados": atualizados, "total": total}


def chave_lotacao(lotacao) -> str:
    """Compara lotacoes sem acento, pontuacao nem espacos ('U.E.B. PROF. SA VALLE' = 'UEB PROF SA VALLE'), mas com
    os numeros ('CRECHE 10' e 'CRECHE 11' sao instituicoes diferentes). E a chave do acesso da instituicao."""
    return re.sub(r"[^A-Z0-9]", "", sem_acento(str(lotacao or "")).upper())


CARGA_PONTUAL = ("turno", "quadro")


def carregar_turnos(con, recs: list[dict], lotacoes: set[str], sem_lotacao: bool) -> dict:
    """Acao pontual: grava o turno e o quadro informados no arquivo (aba CADASTRO) para os servidores das lotacoes
    indicadas (chave_lotacao), substituindo os atuais (campo vazio no arquivo nao apaga). Casa pela matricula; sem matricula no cadastro, pelo CPF ou nome,
    desde que o registro nao seja de outra matricula (outro vinculo da mesma pessoa). sem_lotacao (so RH/admin):
    servidor ainda sem lotacao, achado pela matricula, recebe a lotacao do arquivo."""
    arquivo = {}  # o primeiro registro de cada servidor vale (o CADASTRO vem antes do GEDUC)
    for r in recs:
        d = limpa(r)
        if d.get("nome") and any(d.get(k) for k in CARGA_PONTUAL) and chave_lotacao(d.get("lotacao")) in lotacoes:
            arquivo.setdefault(d.get("matricula") or d.get("cpf") or norm_nome(d["nome"]), d)
    todos = [dict(r) for r in con.execute("SELECT id, nome, matricula, cpf, nome_norm, lotacao, turno, quadro FROM servidores")]
    da_inst = [s for s in todos if chave_lotacao(s["lotacao"]) in lotacoes]
    pela_mat = da_inst + ([s for s in todos if not s["lotacao"]] if sem_lotacao else [])

    usados = set()  # cada servidor do cadastro recebe os dados de uma linha so

    def localiza(d: dict) -> dict | None:
        mat = d.get("matricula")
        if mat and (s := next((s for s in pela_mat if s["matricula"] == mat), None)):
            return s
        for k, v in (("cpf", d.get("cpf")), ("nome_norm", norm_nome(d["nome"]))):
            s = next((s for s in da_inst if v and s[k] == v and s["id"] not in usados
                      and not (mat and s["matricula"])), None)
            if s:
                return s
        return None

    atualizados, iguais, nao_achados, em_outra = 0, 0, [], []
    with con:  # primeiro as linhas com matricula, para a busca por CPF/nome nao tomar o registro delas
        for d in sorted(arquivo.values(), key=lambda d: not d.get("matricula")):
            s = localiza(d)
            if s and s["id"] in usados:
                s = None
            if s:
                usados.add(s["id"])
            if not s:
                outro = any(x for x in todos if (d.get("matricula") and x["matricula"] == d["matricula"])
                            or (d.get("cpf") and x["cpf"] == d["cpf"]))
                (em_outra if outro else nao_achados).append(d["nome"])
            else:
                novo = {k: d[k] for k in CARGA_PONTUAL if d.get(k) and d[k] != s[k]}
                if not s["lotacao"]:
                    novo["lotacao"] = d["lotacao"]
                if not novo:
                    iguais += 1
                    continue
                con.execute(f"UPDATE servidores SET {', '.join(f'{k}=?' for k in novo)}, atualizado_em={_agora()} "
                            "WHERE id=?", (*novo.values(), s["id"]))
                atualizados += 1
    return {"no_arquivo": len(arquivo), "atualizados": atualizados, "iguais": iguais,
            "nao_encontrados": sorted(nao_achados), "outra_lotacao": sorted(em_outra)}


ABAS_SGI = ("CADASTRO", "GEDUC")  # a planilha SGI Servidor tem as duas: le o CADASTRO e completa com o GEDUC
_DIAS_SEMANA = {"SEGUNDA-FEIRA", "TERCA-FEIRA", "QUARTA-FEIRA", "QUINTA-FEIRA", "SEXTA-FEIRA", "SABADO", "DOMINGO"}


def ler_arquivo(nome: str, data: bytes) -> tuple[list[dict], list[tuple[str, str, str]]]:
    """(servidores, feriados) do arquivo: CSV/XLS do GEDUC ou a planilha SGI Servidor (abas CADASTRO, GEDUC e o
    calendario da aba DADOS)."""
    if data[:2] != b"PK":
        return importer.read_file("servidores", nome, data), []
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    abas = {_chave(n): n for n in wb.sheetnames}
    feriados = []
    if "DADOS" in abas:  # colunas G (data), J (rubrica da entrada) e L (rubrica da saida)
        for r in wb[abas["DADOS"]].iter_rows(min_row=3, max_row=800, min_col=7, max_col=12, values_only=True):
            dia, tipo, desc = to_date(r[0]), txt(r[3]), txt(r[5])
            if dia and tipo and _chave(tipo) not in _DIAS_SEMANA:
                tipo = "Ponto facultativo" if "FACULTATIVO" in _chave(tipo) else "Feriado"
                feriados.append((dia, tipo, desc if "FACULTATIVO" not in _chave(desc) else desc or tipo))
    wb.close()
    sgi = [abas[a] for a in ABAS_SGI if a in abas]
    if len(sgi) < 2:
        return importer.read_file("servidores", nome, data), feriados
    recs = []
    for aba in sgi:
        try:
            recs += importer.read_file("servidores", nome, data, sheet=aba)
        except ValueError:
            pass
    return recs, feriados


def _agora() -> str:
    return db.AGORA_PG if db.DATABASE_URL else db.AGORA_SQLITE


# ------------------------------------------------------------------ uso nos documentos

_PARTICULAS = {"DA", "DE", "DO", "DAS", "DOS", "E", "D"}


def abreviado(nome: str) -> str:
    """Etiqueta: os dois primeiros nomes, com a particula do meio ('Ana de Sa' e nao 'Ana De')."""
    partes = (nome or "").split()
    out = []
    for p in partes:
        out.append(p)
        if len([x for x in out if x.upper() not in _PARTICULAS]) == 2:
            break
    return " ".join(x.lower() if x.upper() in _PARTICULAS else x.capitalize() for x in out)


def idade_em(dt_nasc: str, ref: date) -> int | None:
    try:
        n = date.fromisoformat(dt_nasc[:10])
    except (TypeError, ValueError):
        return None
    return ref.year - n.year - ((ref.month, ref.day) < (n.month, n.day))


def mes_ano(ano: int, mes: int) -> str:
    return f"{MESES[mes - 1]} DE {ano}"


# ------------------------------------------------------------------ feriados

def _pascoa(ano: int) -> date:
    """Domingo de Pascoa (algoritmo de Meeus/Butcher)."""
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    return date(ano, mes, (h + l - 7 * m + 114) % 31 + 1)


def feriados_padrao(ano: int) -> list[tuple[str, str, str]]:
    """Feriados nacionais, do Maranhao e de Sao Luis (os mesmos da aba DADOS da planilha). O RH ajusta os pontos
    facultativos de cada ano por decreto."""
    p = _pascoa(ano)
    fixos = [(1, 1, "Feriado", "Confraternização Universal"), (4, 21, "Feriado", "Tiradentes"),
             (5, 1, "Feriado", "Dia do Trabalho"), (6, 29, "Feriado", "São Pedro"),
             (7, 28, "Feriado", "Adesão do Maranhão"), (9, 7, "Feriado", "Independência do Brasil"),
             (9, 8, "Feriado", "Aniversário de São Luís"), (10, 12, "Feriado", "Nossa Senhora Aparecida"),
             (10, 28, "Ponto facultativo", "Dia do Servidor Público"), (11, 2, "Feriado", "Finados"),
             (11, 15, "Feriado", "Proclamação da República"), (11, 20, "Feriado", "Consciência Negra"),
             (12, 8, "Feriado", "Nossa Senhora da Conceição"), (12, 25, "Feriado", "Natal")]
    moveis = [(p - timedelta(days=48), "Feriado", "Carnaval"), (p - timedelta(days=47), "Feriado", "Carnaval"),
              (p - timedelta(days=46), "Ponto facultativo", "Quarta-feira de Cinzas"),
              (p - timedelta(days=3), "Ponto facultativo", "Quinta-feira Santa"),
              (p - timedelta(days=2), "Feriado", "Paixão de Cristo"), (p + timedelta(days=60), "Feriado", "Corpus Christi")]
    out = [(date(ano, m, d).isoformat(), t, n) for m, d, t, n in fixos] + [(d.isoformat(), t, n) for d, t, n in moveis]
    return sorted(out)


def feriados_do_mes(con, ano: int, mes: int) -> dict[int, tuple[str, str]]:
    ini = date(ano, mes, 1)
    r = con.execute("SELECT data, tipo, descricao FROM feriados WHERE data >= ? AND data < ?",
                    (ini.isoformat(), (ini + timedelta(days=32)).replace(day=1).isoformat()))
    return {int(x["data"][8:10]): (x["tipo"], x["descricao"]) for x in r}
